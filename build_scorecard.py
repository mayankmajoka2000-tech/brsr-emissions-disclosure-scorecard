"""
Builds the BRSR Scope 1 & 2 disclosure scorecard workbook.

Inputs : data/brsr_fy26_ghg.csv   (figures extracted from each company's FY2025-26 BRSR)
         ASSESSMENT below         (analyst's level for each rubric criterion, with rationale)
Output : output/BRSR_Scope1_2_Disclosure_Scorecard_FY26.xlsx

Scores in the workbook are live formulas: change a level on the 'Assessment' sheet or a
point value on the 'Rubric' sheet and every total and rank recalculates.

Run: python build_scorecard.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).parent
DATA = pd.read_csv(ROOT / "data" / "brsr_fy26_ghg.csv")
OUT = ROOT / "output" / "BRSR_Scope1_2_Disclosure_Scorecard_FY26.xlsx"

# --------------------------------------------------------------------------- rubric
# (criterion, max, description, [(level, points), ...])
RUBRIC = [
    ("Absolute Scope 1 & 2", 15,
     "Scope 1 and Scope 2 reported in tCO2e for current and prior year at usable precision",
     [("Full", 15), ("Rounded", 12), ("Partial", 7), ("Missing", 0)]),
    ("Gas-wise breakdown", 10,
     "Break-up into CO2, CH4, N2O, HFCs etc., as the BRSR format requests 'if available'",
     [("Full", 10), ("Partial", 5), ("None", 0)]),
    ("Intensity metrics", 15,
     "Per rupee of turnover, PPP-adjusted, and physical-output intensity (5 points each)",
     [("All three", 15), ("Two", 10), ("One", 5), ("None", 0)]),
    ("Scope 2 method", 10,
     "States market- vs location-based; full marks when both figures are given",
     [("Both", 10), ("One stated", 5), ("Not stated", 0)]),
    ("Methodology & factors", 10,
     "Names the standard (GHG Protocol / ISO 14064 / IPCC) and emission-factor sources",
     [("Stated", 10), ("Referenced elsewhere", 5), ("Not stated", 0)]),
    ("External assurance", 15,
     "Independent assurance of the GHG figures and its level",
     [("Reasonable", 15), ("Limited", 8), ("None", 0)]),
    ("Boundary clarity", 5,
     "Clear organisational boundary; consolidated coverage scores higher",
     [("Consolidated", 5), ("Standalone", 3), ("Unclear", 0)]),
    ("Scope 3 (leadership)", 10,
     "Voluntary Scope 3 disclosure; category-level split scores higher",
     [("Category breakdown", 10), ("Total only", 6), ("Not disclosed", 0)]),
    ("Reduction target", 10,
     "Science-based target, or at least a dated net-zero / neutrality commitment",
     [("SBTi", 10), ("Dated net zero", 6), ("None found", 0)]),
]
CRITERIA = [r[0] for r in RUBRIC]

# --------------------------------------------------------------------------- assessment
# level per criterion (order = CRITERIA), plus one-line rationale
ASSESSMENT = {
    "Reliance Industries": (["Full", "Full", "All three", "Both", "Stated", "Reasonable", "Standalone", "Not disclosed", "Dated net zero"],
        "Best-in-class gas split and dual Scope 2; Scope 3 table left without values."),
    "Tata Consultancy Services": (["Full", "Full", "All three", "Both", "Stated", "Reasonable", "Consolidated", "Category breakdown", "SBTi"],
        "Complete on every criterion: gas split, factor sources, 7-category Scope 3, 1.5C SBTi target."),
    "Infosys": (["Full", "None", "All three", "Both", "Stated", "Reasonable", "Standalone", "Total only", "Dated net zero"],
        "Strong, but gas split and Scope 3 categories pushed to the ESG Databook rather than the BRSR."),
    "HDFC Bank": (["Full", "Full", "All three", "One stated", "Stated", "Reasonable", "Standalone", "Total only", "Dated net zero"],
        "Detailed gas split and factors; no location-based Scope 2; financed emissions absent."),
    "ITC": (["Full", "None", "All three", "Both", "Stated", "Reasonable", "Standalone", "Total only", "Dated net zero"],
        "ISO 14064 inventory with dual Scope 2; no gas split; kilotonne precision."),
    "NTPC": (["Full", "None", "All three", "Not stated", "Not stated", "Reasonable", "Standalone", "Total only", "None found"],
        "India's largest emitter: figures assured but no method, Scope 2 basis or reduction target."),
    "UltraTech Cement": (["Full", "None", "All three", "Not stated", "Not stated", "Reasonable", "Standalone", "Total only", "SBTi"],
        "SBTi-validated targets, but BRSR omits methodology and Scope 2 basis."),
    "Larsen & Toubro": (["Full", "None", "Two", "Not stated", "Referenced elsewhere", "Reasonable", "Standalone", "Total only", "Dated net zero"],
        "No physical intensity, gas split or Scope 2 basis; Scope 3 explicitly unassured."),
    "Mahindra & Mahindra": (["Full", "None", "All three", "One stated", "Stated", "Reasonable", "Standalone", "Total only", "SBTi"],
        "GHG Protocol + ISO 14064-3 verified, segment intensities; no gas split."),
    "Tata Steel": (["Rounded", "None", "All three", "One stated", "Stated", "Reasonable", "Consolidated", "Total only", "Dated net zero"],
        "Only company giving standalone and consolidated; figures rounded to whole Mt."),
}

# --------------------------------------------------------------------------- styles
FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1F4E5A")
B = Font(name=FONT, size=10)
BB = Font(name=FONT, size=10, bold=True)
BLUE = Font(name=FONT, size=10, color="0000FF")
GREEN = Font(name=FONT, size=10, color="008000")
TITLE = Font(name=FONT, size=14, bold=True, color="1F4E5A")
SUB = Font(name=FONT, size=10, italic=True, color="555555")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
thin = Side(style="thin", color="D0D7DE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def header(ws, row, labels, widths=None):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill, cell.alignment, cell.border = H, HFILL, CENTER, BOX
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 32


def body(cell, font=B, fmt=None, align=WRAP):
    cell.font, cell.border, cell.alignment = font, BOX, align
    if fmt:
        cell.number_format = fmt


wb = Workbook()

# --------------------------------------------------------------------------- About
ws = wb.active
ws.title = "About"
ws.column_dimensions["A"].width = 110
lines = [
    ("BRSR Scope 1 & 2 Disclosure Scorecard: 10 Nifty 50 companies, FY2025-26", TITLE),
    ("Quality of greenhouse-gas disclosure in SEBI Business Responsibility & Sustainability Reports", SUB),
    ("", B),
    ("What this scores", BB),
    ("How well each company discloses its Scope 1 and 2 emissions in its FY2025-26 BRSR (Principle 6, Essential Indicator 7, "
     "plus the Scope 3 leadership indicator). It scores disclosure quality and transparency, not emissions performance: "
     "a heavy emitter can score well if it reports clearly.", B),
    ("", B),
    ("Sheets", BB),
    ("Scores: total score out of 100, rank and grade, computed by formula from the Assessment and Rubric sheets.", B),
    ("Assessment: the level assigned to each company on each criterion (yellow cells, blue text = editable input), with rationale.", B),
    ("Rubric: the 9 criteria, their weights and the points for each level. Edit points here and all scores update.", B),
    ("Emissions Data: Scope 1, 2 and 3 figures as reported, year-on-year change, intensities, assurer and target, with source links.", B),
    ("", B),
    ("Method", BB),
    ("1. Sample: 10 Nifty 50 constituents across 9 sectors, weighted toward high emitters (power, cement, steel, oil) "
     "with IT, banking, FMCG, autos and engineering for contrast.", B),
    ("2. Source: each company's FY2025-26 BRSR as filed with NSE or published on its website (filed May-Aug 2026). "
     "Figures were extracted from the PDF text layer and checked against the GHG table of each report.", B),
    ("3. Scoring: 9 criteria weighted to 100 points, built from the BRSR format itself (SEBI circular, BRSR Core) "
     "and from what assurance and investor frameworks treat as decision-useful (GHG Protocol Scope 2 guidance, ISO 14064, SBTi).", B),
    ("4. Grades: A >= 85, B 70-84, C 55-69, D < 55.", B),
    ("", B),
    ("Limitations", BB),
    ("Only the BRSR document is scored. Several companies disclose more in separate ESG databooks or climate reports "
     "(e.g. Infosys gas split, Tata Steel Climate Change Report), which this scorecard deliberately does not credit, "
     "since BRSR is the regulated disclosure. Levels are analyst judgements and are documented on the Assessment sheet.", B),
    ("Reported units differ (tCO2e, kilotonnes, million tonnes); the Emissions Data sheet converts all to tCO2e.", B),
    ("", B),
    ("Prepared by Mayank Majoka, October 2026. Not investment advice.", SUB),
]
for i, (text, font) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=text)
    c.font, c.alignment = font, Alignment(wrap_text=True, vertical="top")

# --------------------------------------------------------------------------- Rubric
rb = wb.create_sheet("Rubric")
header(rb, 1, ["Criterion", "Level", "Points", "Max points", "What it measures"], [26, 24, 10, 12, 80])
r = 2
for crit, mx, desc, levels in RUBRIC:
    for j, (lvl, pts) in enumerate(levels):
        vals = [crit, lvl, pts, mx if j == 0 else None, desc if j == 0 else None]
        for c, v in enumerate(vals, 1):
            cell = rb.cell(row=r, column=c, value=v)
            body(cell, BLUE if c == 3 else B, align=CENTER if c in (3, 4) else WRAP)
        r += 1
RUBRIC_LAST = r - 1
rb.cell(row=r + 1, column=1, value="Total possible").font = BB
rb.cell(row=r + 1, column=4, value=f"=SUM(D2:D{RUBRIC_LAST})").font = BB
rb.cell(row=r + 2, column=1, value="Blue values are editable inputs; scores on the Scores sheet update automatically.").font = SUB
rb.freeze_panes = "A2"

# --------------------------------------------------------------------------- Assessment
asm = wb.create_sheet("Assessment")
header(asm, 1, ["Company"] + CRITERIA + ["Rationale"], [26] + [15] * len(CRITERIA) + [70])
companies = list(DATA["company"])
for i, comp in enumerate(companies, 2):
    levels, why = ASSESSMENT[comp]
    body(asm.cell(row=i, column=1, value=comp), BB)
    for j, lvl in enumerate(levels, 2):
        c = asm.cell(row=i, column=j, value=lvl)
        body(c, BLUE, align=CENTER)
        c.fill = INPUT_FILL
    body(asm.cell(row=i, column=len(CRITERIA) + 2, value=why))
    asm.row_dimensions[i].height = 30
# dropdowns restricted to each criterion's levels
for j, (crit, _, _, levels) in enumerate(RUBRIC, 2):
    col = get_column_letter(j)
    dv = DataValidation(type="list", formula1='"' + ",".join(l for l, _ in levels) + '"', allow_blank=False)
    dv.add(f"{col}2:{col}{len(companies) + 1}")
    asm.add_data_validation(dv)
asm.cell(row=len(companies) + 3, column=1,
         value="Yellow cells are inputs: choose a level from the dropdown. Points come from the Rubric sheet.").font = SUB
asm.freeze_panes = "B2"

# --------------------------------------------------------------------------- Scores
sc = wb.create_sheet("Scores", 1)
sc["A1"] = "BRSR Scope 1 & 2 Disclosure Scores, FY2025-26"
sc["A1"].font = TITLE
sc["A2"] = "Out of 100. Higher = more complete, transparent and assured disclosure (not lower emissions)."
sc["A2"].font = SUB
hdr = ["Company", "Sector"] + CRITERIA + ["Total (/100)", "Rank", "Grade"]
header(sc, 4, hdr, [26, 22] + [12] * len(CRITERIA) + [12, 8, 8])
n_crit = len(CRITERIA)
first, last = 5, 4 + len(companies)
for i, comp in enumerate(companies):
    row = first + i
    arow = i + 2
    body(sc.cell(row=row, column=1, value=f"=Assessment!A{arow}"), Font(name=FONT, size=10, bold=True, color="008000"))
    body(sc.cell(row=row, column=2, value=DATA.loc[i, "sector"]))
    for j, crit in enumerate(CRITERIA):
        acol = get_column_letter(j + 2)
        f = (f'=SUMIFS(Rubric!$C$2:$C${RUBRIC_LAST},Rubric!$A$2:$A${RUBRIC_LAST},"{crit}",'
             f'Rubric!$B$2:$B${RUBRIC_LAST},Assessment!{acol}{arow})')
        body(sc.cell(row=row, column=j + 3, value=f), align=CENTER)
    tcol = n_crit + 3
    s, e = get_column_letter(3), get_column_letter(n_crit + 2)
    T = get_column_letter(tcol)
    body(sc.cell(row=row, column=tcol, value=f"=SUM({s}{row}:{e}{row})"), BB, align=CENTER)
    body(sc.cell(row=row, column=tcol + 1, value=f"=RANK({T}{row},${T}${first}:${T}${last})"), align=CENTER)
    body(sc.cell(row=row, column=tcol + 2,
                 value=f'=IF({T}{row}>=85,"A",IF({T}{row}>=70,"B",IF({T}{row}>=55,"C","D")))'), BB, align=CENTER)

# averages + max row
avg_row = last + 1
body(sc.cell(row=avg_row, column=1, value="Average"), BB)
body(sc.cell(row=avg_row, column=2, value=""))
max_row = avg_row + 1
body(sc.cell(row=max_row, column=1, value="Maximum possible"), BB)
body(sc.cell(row=max_row, column=2, value=""))
for j, crit in enumerate(CRITERIA + ["Total"]):
    col = j + 3
    L = get_column_letter(col)
    body(sc.cell(row=avg_row, column=col, value=f"=AVERAGE({L}{first}:{L}{last})"), BB, "0.0", CENTER)
    if crit == "Total":
        mx = f"=Rubric!D{RUBRIC_LAST + 2}"
    else:
        mx = f'=SUMIFS(Rubric!$D$2:$D${RUBRIC_LAST},Rubric!$A$2:$A${RUBRIC_LAST},"{crit}")'
    body(sc.cell(row=max_row, column=col, value=mx), align=CENTER)
# % of max achieved on average - shows which criteria are weakest market-wide
pct_row = max_row + 1
body(sc.cell(row=pct_row, column=1, value="Average as % of max"), BB)
body(sc.cell(row=pct_row, column=2, value=""))
for j in range(n_crit + 1):
    L = get_column_letter(j + 3)
    body(sc.cell(row=pct_row, column=j + 3, value=f"=IF({L}{max_row}=0,0,{L}{avg_row}/{L}{max_row})"), BB, "0%", CENTER)

T = get_column_letter(n_crit + 3)
sc.conditional_formatting.add(f"{T}{first}:{T}{last}", ColorScaleRule(
    start_type="num", start_value=40, start_color="F8D7DA",
    mid_type="num", mid_value=70, mid_color="FFF3CD",
    end_type="num", end_value=100, end_color="D4EDDA"))
sc.conditional_formatting.add(f"C{pct_row}:{get_column_letter(n_crit + 2)}{pct_row}", ColorScaleRule(
    start_type="num", start_value=0, start_color="F8D7DA",
    end_type="num", end_value=1, end_color="D4EDDA"))
sc.freeze_panes = "C5"

chart = BarChart()
chart.type = "bar"
chart.title = "Disclosure score (out of 100)"
chart.y_axis.title = "Score"
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 100
chart.legend = None
chart.add_data(Reference(sc, min_col=n_crit + 3, min_row=4, max_row=last), titles_from_data=True)
chart.set_categories(Reference(sc, min_col=1, min_row=first, max_row=last))
chart.height, chart.width = 9, 18
sc.add_chart(chart, f"A{pct_row + 3}")

# --------------------------------------------------------------------------- Emissions Data
ed = wb.create_sheet("Emissions Data")
cols = ["Company", "Sector", "Boundary", "Scope 1 FY26 (tCO2e)", "Scope 1 FY25 (tCO2e)", "Scope 1 YoY",
        "Scope 2 FY26 (tCO2e)", "Scope 2 FY25 (tCO2e)", "Scope 2 YoY", "Scope 1+2 FY26 (tCO2e)", "Scope 1 share",
        "Scope 2 basis reported", "Scope 2 location-based FY26", "Scope 3 FY26 (tCO2e)", "Scope 3 FY25 (tCO2e)",
        "Physical intensity FY26", "Intensity unit", "Assurer", "Assurance level", "Stated target", "Notes", "Source (BRSR FY2025-26)"]
header(ed, 1, cols, [24, 18, 14, 15, 15, 9, 14, 14, 9, 15, 9, 18, 15, 15, 15, 12, 26, 26, 12, 40, 60, 30])
num = "#,##0;(#,##0);-"
pct = "0.0%;-0.0%;-"
for i, rec in DATA.iterrows():
    r = i + 2
    vals = {
        1: rec.company, 2: rec.sector, 3: rec.boundary_basis,
        4: rec.scope1_fy26_tco2e, 5: rec.scope1_fy25_tco2e,
        6: f"=IF(E{r}=0,\"\",D{r}/E{r}-1)",
        7: rec.scope2_fy26_tco2e, 8: rec.scope2_fy25_tco2e,
        9: f"=IF(H{r}=0,\"\",G{r}/H{r}-1)",
        10: f"=D{r}+G{r}", 11: f"=IF(J{r}=0,\"\",D{r}/J{r})",
        12: rec.scope2_basis_reported,
        13: None if pd.isna(rec.scope2_location_fy26_tco2e) else rec.scope2_location_fy26_tco2e,
        14: None if pd.isna(rec.scope3_fy26_tco2e) else rec.scope3_fy26_tco2e,
        15: None if pd.isna(rec.scope3_fy25_tco2e) else rec.scope3_fy25_tco2e,
        16: None if pd.isna(rec.physical_intensity_fy26) else rec.physical_intensity_fy26,
        17: "" if pd.isna(rec.physical_intensity_unit) else rec.physical_intensity_unit,
        18: rec.assurance_provider, 19: rec.assurance_level, 20: rec.stated_target, 21: rec.notes,
    }
    for c, v in vals.items():
        cell = ed.cell(row=r, column=c, value=v)
        is_input = c in (4, 5, 7, 8, 13, 14, 15, 16)
        fmt = num if c in (4, 5, 7, 8, 10, 13, 14, 15) else pct if c in (6, 9, 11) else ("0.000" if c == 16 else None)
        body(cell, BLUE if is_input else B, fmt, CENTER if c in (6, 9, 11, 16, 19) else WRAP)
    link = ed.cell(row=r, column=22, value="Open report")
    link.hyperlink = rec.report_url
    body(link, Font(name=FONT, size=10, color="0563C1", underline="single"))
    ed.row_dimensions[r].height = 64
for c, unit_note in ((4, "ITC reports in kilotonnes and Tata Steel in million tonnes; converted to tCO2e here."),):
    ed.cell(row=1, column=c).comment = Comment(unit_note, "Analyst")
tot = len(DATA) + 2
body(ed.cell(row=tot, column=1, value="Total (10 companies)"), BB)
for c in (4, 5, 7, 8, 10):
    L = get_column_letter(c)
    body(ed.cell(row=tot, column=c, value=f"=SUM({L}2:{L}{tot - 1})"), BB, num)
body(ed.cell(row=tot, column=6, value=f"=D{tot}/E{tot}-1"), BB, pct, CENTER)
body(ed.cell(row=tot, column=9, value=f"=G{tot}/H{tot}-1"), BB, pct, CENTER)
ed.cell(row=tot + 2, column=1, value="Blue figures are as reported in each FY2025-26 BRSR (Principle 6). "
        "Black cells are formulas.").font = SUB
ed.freeze_panes = "B2"

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"Saved {OUT}")
